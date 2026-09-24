"""Seleção, mesclagem e validação das planilhas de índices."""
from __future__ import annotations
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable
import pandas as pd
from openpyxl import Workbook, load_workbook
from judicial_calc.data_sources.local_excel import (
    DIARIA_12_6_XLSX, DIARIA_SELIC_IPCAE_XLSX, MENSAL_XLSX, RATE_DECIMAL, VALUE_INDEX,
    local_index_specs,
)
from .models import DrCalcSeries, DrCalcUpdateError
from .parsing import _competencia_from_period, _normalize_text

def _series_score(series_name: str, target_label: str, extra_aliases: Iterable[str] = ()) -> int:
    """Calcula uma pontuação heurística para escolher a série mais compatível com um destino."""
    series_norm = _normalize_text(series_name)
    target_norm = _normalize_text(target_label)
    aliases = {_normalize_text(target_label), *[_normalize_text(a) for a in extra_aliases if a]}
    score = 0
    for alias in aliases:
        if not alias:
            continue
        if alias == series_norm:
            score = max(score, 100)
        elif alias in series_norm or series_norm in alias:
            score = max(score, 80)
        else:
            tokens = [t for t in alias.split("_") if len(t) >= 3]
            if tokens:
                common = sum(1 for t in tokens if t in series_norm)
                score = max(score, int(60 * common / len(tokens)))
    return score


def _looks_like_rate_percent(series: DrCalcSeries) -> bool:
    """Avalia de forma conservadora se uma série sem unidade parece percentual.

    A heurística só é usada quando o HTML não informa claramente se o valor é
    variação ou número-índice. Evita gravar, por exemplo, um número-índice de
    centenas como se fosse uma taxa mensal em decimal.
    """
    values = sorted(abs(record.valor) for record in series.records if record.valor is not None)
    if not values:
        return False
    median = values[len(values) // 2]
    p95 = values[min(len(values) - 1, int(len(values) * Decimal("0.95")))] if len(values) > 1 else values[0]
    return median <= Decimal("10") and p95 <= Decimal("50")


def _best_series_for_monthly_column(column: str, series_list: list[DrCalcSeries]) -> DrCalcSeries | None:
    """Seleciona a melhor série mensal para alimentar uma coluna da planilha local."""
    spec = None
    for s in local_index_specs():
        if s.column.strip() == column.strip():
            spec = s
            break
    aliases = []
    if spec:
        aliases.extend([spec.key, spec.column, *spec.aliases])
    best: tuple[int, DrCalcSeries] | None = None
    for series in series_list:
        if series.periodicity != "mensal":
            continue
        score = _series_score(series.name, column, aliases)
        if spec is not None:
            expected_metric = "rate_percent" if spec.mode == RATE_DECIMAL else "index_value"
            if series.metric == expected_metric:
                score += 25
            elif series.metric != "unknown":
                score -= 35
            elif spec.mode == VALUE_INDEX:
                # Sem unidade explícita não é seguro substituir número-índice.
                continue
            elif not _looks_like_rate_percent(series):
                continue
        if score >= 55 and (best is None or score > best[0]):
            best = (score, series)
    return best[1] if best else None


def _best_daily_series(kind: str, series_list: list[DrCalcSeries]) -> DrCalcSeries | None:
    """Seleciona a melhor série diária entre as séries descobertas."""
    if kind == "selic_ipcae":
        aliases = [
            "TAXA LEGAL DIARIA SELIC IPCAE",
            "TAXA LEGAL DIÁRIA SELIC IPCA-E",
            "SELIC IPCA-E",
            "SELIC-IPCAE",
            "Lei 14905 STJ Tema 1368 SELIC IPCA",
        ]
    else:
        aliases = [
            "TAXA LEGAL 12% aa 6% aa",
            "TAXA LEGAL 12 aa 6 aa",
            "12% a.a. 6% a.a.",
            "art. 406 CC 12% 6%",
        ]
    best: tuple[int, DrCalcSeries] | None = None
    for series in series_list:
        if series.periodicity != "diaria":
            continue
        score = max(_series_score(series.name, alias, aliases) for alias in aliases)
        if series.metric == "rate_percent":
            score += 15
        elif series.metric == "index_value":
            score -= 30
        if kind == "selic_ipcae":
            norm = _normalize_text(series.name)
            if "selic" in norm and ("ipcae" in norm or "ipca_e" in norm or "ipca" in norm):
                score += 20
        else:
            norm = _normalize_text(series.name)
            if "12" in norm and "6" in norm:
                score += 20
        if score >= 65 and (best is None or score > best[0]):
            best = (score, series)
    return best[1] if best else None


def _convert_monthly_value(column: str, valor: Decimal, *, source_metric: str = "unknown") -> Decimal:
    """Converte o valor mensal da fonte para a unidade esperada pela planilha local.

    Quando o parser identificou explicitamente ``rate_percent``, o valor do
    DrCalc está em percentual e deve ser dividido por 100 inclusive para taxas
    pequenas, como ``0,03%``. O limiar numérico fica somente para fontes legadas
    sem metadado de unidade.
    """
    spec = None
    for s in local_index_specs():
        if s.column.strip() == column.strip():
            spec = s
            break
    if spec and spec.mode == RATE_DECIMAL:
        if source_metric == "rate_percent":
            return valor / Decimal("100")
        if source_metric == "unknown" and abs(valor) > Decimal("0.05"):
            # Compatibilidade com o fallback legado, cuja unidade nem sempre é
            # identificável pelo cabeçalho da tabela.
            return valor / Decimal("100")
    return valor


def _convert_daily_value(valor: Decimal, *, source_metric: str = "unknown") -> Decimal:
    # Taxa diária em percentual, como 0,03, vira decimal 0,0003.
    """Converte o valor diário da fonte para a unidade esperada pela planilha local."""
    if source_metric == "rate_percent":
        return valor / Decimal("100")
    if source_metric == "unknown" and abs(valor) > Decimal("0.01"):
        return valor / Decimal("100")
    return valor


def _monthly_dataframe_from_workbook(path: Path) -> pd.DataFrame:
    """Lê a planilha mensal existente e normaliza seu conteúdo em DataFrame."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["indices"] if "indices" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise DrCalcUpdateError(f"Planilha mensal vazia: {path}")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    data = []
    for raw in rows[1:]:
        if not raw or raw[0] is None:
            continue
        row = {headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))}
        data.append(row)
    return pd.DataFrame(data, columns=headers)


def _daily_dataframe_from_workbook(path: Path) -> pd.DataFrame:
    """Lê a planilha diária existente e normaliza seu conteúdo em DataFrame."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise DrCalcUpdateError(f"Planilha diária vazia: {path}")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    return pd.DataFrame([{headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))} for raw in rows[1:] if raw and raw[0] is not None])


def _write_monthly_workbook(df: pd.DataFrame, path: Path) -> None:
    """Grava a tabela mensal normalizada no arquivo Excel de destino."""
    wb = Workbook()
    ws = wb.active
    ws.title = "indices"
    for c_idx, header in enumerate(df.columns, start=1):
        ws.cell(row=1, column=c_idx, value=header)
    for r_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for c_idx, column in enumerate(df.columns, start=1):
            value = row[column]
            if pd.isna(value):
                value = None
            if isinstance(value, Decimal):
                value = float(value)
            ws.cell(row=r_idx, column=c_idx, value=value)
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 12
    for col in range(2, min(len(df.columns), 60) + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 18
    wb.save(path)


def _write_daily_workbook(df: pd.DataFrame, path: Path) -> None:
    """Grava a tabela diária normalizada no arquivo Excel de destino."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["data", "valor_indice"])
    for _, row in df.iterrows():
        data_value = row["data"]
        if isinstance(data_value, pd.Timestamp):
            data_value = data_value.date()
        valor = row["valor_indice"]
        if isinstance(valor, Decimal):
            valor = float(valor)
        ws.append([data_value, valor])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    wb.save(path)


def _merge_monthly(existing_path: Path, series_list: list[DrCalcSeries]) -> tuple[pd.DataFrame, list[str], list[str]]:
    """Combina séries mensais baixadas com a estrutura da planilha local."""
    df = _monthly_dataframe_from_workbook(existing_path)
    if "data" not in df.columns:
        raise DrCalcUpdateError("taxas_mensais.xlsx precisa ter coluna 'data'.")
    df["data"] = df["data"].astype(str).str[:7]
    df = df.drop_duplicates("data", keep="last").sort_values("data").reset_index(drop=True)
    df = df.set_index("data")

    updated: list[str] = []
    preserved: list[str] = []
    for column in list(df.columns):
        series = _best_series_for_monthly_column(column, series_list)
        if series is None:
            preserved.append(column)
            continue
        changed = False
        for record in series.records:
            competencia = _competencia_from_period(record.periodo)
            value = _convert_monthly_value(column, record.valor, source_metric=series.metric)
            if competencia not in df.index:
                df.loc[competencia, :] = pd.NA
            df.loc[competencia, column] = float(value)
            changed = True
        if changed:
            updated.append(f"{column} <= {series.name}")
        else:
            preserved.append(column)

    if not updated:
        raise DrCalcUpdateError("Nenhuma série mensal compatível foi encontrada no DrCalc.")

    df = df.sort_index().reset_index().rename(columns={"index": "data"})
    ordered = ["data"] + [c for c in _monthly_dataframe_from_workbook(existing_path).columns if c != "data"]
    return df[ordered], updated, preserved


def _merge_daily(existing_path: Path, series: DrCalcSeries | None, label: str) -> tuple[pd.DataFrame, list[str]]:
    """Combina a série diária baixada com a estrutura da planilha local."""
    df = _daily_dataframe_from_workbook(existing_path)
    if "data" not in df.columns or "valor_indice" not in df.columns:
        raise DrCalcUpdateError(f"{existing_path.name} precisa ter colunas data e valor_indice.")
    df["data"] = pd.to_datetime(df["data"]).dt.date
    df = df.drop_duplicates("data", keep="last").sort_values("data").reset_index(drop=True)
    if series is None:
        raise DrCalcUpdateError(f"Série diária não encontrada no DrCalc: {label}.")

    by_date = {row["data"]: row["valor_indice"] for _, row in df.iterrows()}
    for record in series.records:
        if not isinstance(record.periodo, date):
            continue
        by_date[record.periodo] = float(_convert_daily_value(record.valor, source_metric=series.metric))
    if not by_date:
        raise DrCalcUpdateError(f"Nenhum dado diário válido para {label}.")
    out = pd.DataFrame([{"data": k, "valor_indice": v} for k, v in sorted(by_date.items())])
    return out, [f"{label} <= {series.name}"]


def _validate_outputs(monthly: pd.DataFrame, daily_selic_ipcae: pd.DataFrame, daily_12_6: pd.DataFrame) -> None:
    """Valida se as planilhas produzidas possuem estrutura e conteúdo mínimos esperados."""
    if monthly.empty or "data" not in monthly.columns:
        raise DrCalcUpdateError("Saída mensal inválida.")
    if len(monthly.columns) < 5:
        raise DrCalcUpdateError("Saída mensal perdeu colunas demais.")
    for label, df in ((DIARIA_SELIC_IPCAE_XLSX, daily_selic_ipcae), (DIARIA_12_6_XLSX, daily_12_6)):
        if df.empty or set(["data", "valor_indice"]) - set(df.columns):
            raise DrCalcUpdateError(f"Saída diária inválida: {label}.")
        if df["valor_indice"].isna().any():
            raise DrCalcUpdateError(f"Saída diária contém valores vazios: {label}.")

def _df_row_count(path: Path, *, daily: bool = False) -> int:
    """Conta linhas úteis de uma planilha local de índices."""
    if not path.exists():
        return 0
    try:
        df = _daily_dataframe_from_workbook(path) if daily else _monthly_dataframe_from_workbook(path)
        return int(len(df))
    except Exception:
        return 0


def _df_date_range(df: pd.DataFrame, date_col: str = "data") -> tuple[str, str]:
    """Retorna intervalo textual mínimo/máximo de datas/competências."""
    if df.empty or date_col not in df.columns:
        return "", ""
    values = df[date_col].dropna().astype(str)
    if values.empty:
        return "", ""
    return str(values.min()), str(values.max())


def _build_diff_row(filename: str, before_rows: int, after_df: pd.DataFrame) -> dict[str, Any]:
    """Resume diferença de linhas e intervalo de datas de uma planilha."""
    date_min, date_max = _df_date_range(after_df)
    after_rows = int(len(after_df))
    return {
        "arquivo": filename,
        "linhas_antes": int(before_rows),
        "linhas_depois": after_rows,
        "linhas_adicionadas_liquidas": after_rows - int(before_rows),
        "data_minima": date_min,
        "data_maxima": date_max,
    }


def _monthly_last_competence_by_column(df: pd.DataFrame) -> dict[str, str]:
    """Obtém a última competência não nula de cada série mensal."""
    if df.empty or "data" not in df.columns:
        return {}
    dates = df["data"].astype(str).str[:7]
    result: dict[str, str] = {}
    for column in df.columns:
        if column == "data":
            continue
        mask = df[column].notna()
        if mask.any():
            result[str(column)] = str(dates[mask].max())
    return result


def _daily_last_competence(df: pd.DataFrame) -> str:
    """Retorna a última data útil de uma série diária como texto ISO."""
    if df.empty or "data" not in df.columns:
        return ""
    values = pd.to_datetime(df["data"], errors="coerce").dropna()
    if values.empty:
        return ""
    return values.max().date().isoformat()


def _coverage_advancements(
    before_monthly: pd.DataFrame,
    after_monthly: pd.DataFrame,
    before_daily_selic: pd.DataFrame,
    after_daily_selic: pd.DataFrame,
    before_daily_12_6: pd.DataFrame,
    after_daily_12_6: pd.DataFrame,
) -> list[dict[str, str]]:
    """Lista somente séries cuja competência máxima realmente avançou."""
    advancements: list[dict[str, str]] = []
    before_month = _monthly_last_competence_by_column(before_monthly)
    after_month = _monthly_last_competence_by_column(after_monthly)
    for column, after in sorted(after_month.items()):
        before = before_month.get(column, "")
        if after and (not before or after > before):
            advancements.append({"arquivo": MENSAL_XLSX, "serie": column, "antes": before, "depois": after})

    for filename, label, before_df, after_df in (
        (DIARIA_SELIC_IPCAE_XLSX, "TAXA LEGAL DIARIA (SELIC-IPCAE)", before_daily_selic, after_daily_selic),
        (DIARIA_12_6_XLSX, "TAXA LEGAL - 12% aa - 6% aa", before_daily_12_6, after_daily_12_6),
    ):
        before = _daily_last_competence(before_df)
        after = _daily_last_competence(after_df)
        if after and (not before or after > before):
            advancements.append({"arquivo": filename, "serie": label, "antes": before, "depois": after})
    return advancements


def _consistency_checks(monthly: pd.DataFrame, daily_selic_ipcae: pd.DataFrame, daily_12_6: pd.DataFrame) -> list[dict[str, Any]]:
    """Gera checagens legíveis para a tela administrativa de índices."""
    checks = [
        {
            "checagem": "taxas_mensais.xlsx contém coluna data",
            "status": "ok" if "data" in monthly.columns else "falha",
            "detalhe": f"{len(monthly)} linhas; {len(monthly.columns)} colunas",
        },
        {
            "checagem": "taxas_mensais.xlsx preserva ao menos 5 colunas",
            "status": "ok" if len(monthly.columns) >= 5 else "falha",
            "detalhe": f"{len(monthly.columns)} colunas encontradas",
        },
        {
            "checagem": f"{DIARIA_SELIC_IPCAE_XLSX} contém data e valor_indice",
            "status": "ok" if {"data", "valor_indice"}.issubset(daily_selic_ipcae.columns) else "falha",
            "detalhe": f"{len(daily_selic_ipcae)} linhas",
        },
        {
            "checagem": f"{DIARIA_12_6_XLSX} contém data e valor_indice",
            "status": "ok" if {"data", "valor_indice"}.issubset(daily_12_6.columns) else "falha",
            "detalhe": f"{len(daily_12_6)} linhas",
        },
        {
            "checagem": "Planilhas diárias sem valores vazios",
            "status": "ok" if not daily_selic_ipcae.get("valor_indice", pd.Series(dtype=float)).isna().any() and not daily_12_6.get("valor_indice", pd.Series(dtype=float)).isna().any() else "falha",
            "detalhe": "Coluna valor_indice validada antes da substituição atômica.",
        },
    ]
    return checks
