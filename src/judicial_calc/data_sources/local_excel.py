"""Leitura das planilhas locais usadas nos cálculos judiciais.

As planilhas são distribuídas dentro de ``judicial_calc.data`` para que os
cálculos possam ser reproduzidos offline quando a fonte oficial já estiver
materializada no projeto.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from functools import lru_cache
from importlib import resources
from pathlib import Path
from typing import Any, Literal
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from judicial_calc.core.dates import somar_meses
from judicial_calc.core.numbers import D

RATE_DECIMAL: Literal["rate_decimal"] = "rate_decimal"
VALUE_INDEX: Literal["value_index"] = "value_index"
IndexMode = Literal["rate_decimal", "value_index"]

MENSAL_XLSX = "taxas_mensais.xlsx"
DIARIA_SELIC_IPCAE_XLSX = "TAXA LEGAL DIARIA (SELIC-IPCAE).xlsx"
DIARIA_12_6_XLSX = "TAXA LEGAL - 12% aa - 6% aa.xlsx"


@dataclass(frozen=True)
class LocalIndexSpec:
    """Metadados de uma coluna de índice na tabela mensal local.

    ``mode`` define como a coluna deve ser acumulada:
    - ``rate_decimal``: a célula é variação mensal em decimal. Ex.: 0.0062 = 0,62%.
    - ``value_index``: a célula é número-índice/valor acumulado. O fator é final/inicial.
    """

    key: str
    column: str
    mode: IndexMode
    aliases: tuple[str, ...] = ()


@dataclass(frozen=True)
class MissingLocalIndexSpec:
    """Metadados de índice conhecido, mas sem coluna na planilha local.

    Entrada de construção:
        ``key``: chave técnica; ``label``: nome exibível;
        ``available_range``: período conhecido em documentação externa.

    Saída:
        Instância imutável usada para mensagens de erro claras.
    """

    key: str
    label: str
    available_range: str


@dataclass(frozen=True)
class LocalIndexCoverage:
    """Cobertura efetivamente observada para uma coluna da planilha mensal.

    ``last_competence`` representa a última competência com valor não nulo no
    arquivo local. ``maximum_update_competence`` representa o maior mês de
    atualização que o motor consegue calcular com essa série. Para índices de
    variação mensal, o cálculo do mês M usa a taxa de M-1; para números-índice,
    o próprio mês M precisa existir.
    """

    key: str
    label: str
    mode: IndexMode
    first_competence: str
    last_competence: str
    maximum_update_competence: str


def _resource_path(filename: str) -> Path:
    """Resolve o caminho de uma planilha empacotada no projeto.

    Entrada:
        ``filename``: nome do arquivo em ``src/judicial_calc/data``.

    Saída:
        ``Path`` absoluto para leitura pelo ``openpyxl``.

    Exemplo:
        ``_resource_path("taxas_mensais.xlsx")`` aponta para a planilha mensal.
    """
    return Path(resources.files("judicial_calc.data").joinpath(filename))


def normalize_key(text: str) -> str:
    """Normaliza nomes e rótulos de índices para chaves seguras de API.

    Entrada:
        ``text``: texto livre, por exemplo ``"IPCA-15 (IBGE)"``.

    Saída:
        ``str`` em minúsculas, sem acentos e com separador ``_``.
        Ex.: ``"ipca_15_ibge"``.
    """
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode("ascii")
    text = text.lower().replace("%", " pct ")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return re.sub(r"_+", "_", text).strip("_")


def _spec(key: str, column: str, mode: IndexMode, *aliases: str) -> LocalIndexSpec:
    """Cria a especificação de uma coluna mensal com aliases normalizados.

    Entrada:
        ``key``: chave pública; ``column``: nome da coluna na planilha;
        ``mode``: ``rate_decimal`` ou ``value_index``;
        ``aliases``: nomes alternativos aceitos.

    Saída:
        ``LocalIndexSpec`` pronto para registro.
    """
    auto_aliases = {normalize_key(column), normalize_key(key), *[normalize_key(a) for a in aliases]}
    auto_aliases.discard(key)
    return LocalIndexSpec(key=key, column=column.strip(), mode=mode, aliases=tuple(sorted(auto_aliases)))


# Metadados das colunas disponíveis no arquivo mensal anexado. A lista cobre os
# índices visíveis nas imagens do usuário quando eles estão presentes na planilha.
_INDEX_SPECS: tuple[LocalIndexSpec, ...] = (
    _spec("cesta_basica_sao_paulo", "Cesta Básica", VALUE_INDEX, "Cesta básica (São Paulo)", "cesta_basica"),
    _spec("cub_sinduscon_sp", "CUB-SINDUSCON/SP", RATE_DECIMAL, "CUB-SINDUSCON/SP"),
    _spec("icv_dieese", "ICV-DIEESE", RATE_DECIMAL),
    _spec("igp_di_fgv", "IGP-DI (FGV)", RATE_DECIMAL),
    _spec("igp_m_fgv", "IGP-M - (FGV)", RATE_DECIMAL, "IGP-M (FGV)", "igpm_fgv"),
    _spec("incc_di_fgv", "INCC-DI - (FGV)", RATE_DECIMAL),
    _spec("inpc_ibge", "INPC-IBGE", RATE_DECIMAL, "INPC (IBGE)"),
    _spec("ipa_di_fgv", "IPA-DI - (FGV)", RATE_DECIMAL),
    _spec("ipca_ibge", "IPCA (IBGE)", RATE_DECIMAL),
    _spec("ipca_15_ibge", "IPCA-15 (IBGE)", RATE_DECIMAL),
    _spec("ipca_e_ibge", "IPCA-E (IBGE)", RATE_DECIMAL, "IPCA-E IBGE"),
    _spec("ipc_di_fgv", "IPC-DI - (FGV)", RATE_DECIMAL),
    _spec("ipc_fipe", "IPC-FIPE", RATE_DECIMAL),
    _spec("ist_telecomunicacoes", "IST - Telecomunicações", RATE_DECIMAL),
    _spec("salario_minimo", "Salário Mínimo", VALUE_INDEX),
    _spec("debitos_judiciais_acoes_acidentarias", "Débitos Judiciais relativos às Ações Acidentárias", VALUE_INDEX, "Ações Acidentárias"),
    _spec("encoge_xi_encontro", "ENCOGE (XI ENCONTRO)", VALUE_INDEX, "ENCOGE"),
    _spec("jf_beneficio_previdenciario_res_267_2013", "JF-Benefício Previdenciário (Res.267/2013)", RATE_DECIMAL),
    _spec("jf_condenatorias_fazenda_publica", "JF-Condenatórias da Fazenda Pública", VALUE_INDEX),
    _spec("jf_condenatorias_geral_exceto_fazenda_publica", "JF-Condenatórias Geral (exceto Fazenda Pública)", RATE_DECIMAL),
    _spec("jf_desapropriacoes_res_267_2013", "JF-Desapropriações (Res.267/2013)", RATE_DECIMAL),
    _spec("precatorios_acoes_acidentarias_ec_62_2009", "Precatórios e Ações Acidentárias - EC 62/2009", VALUE_INDEX),
    _spec("tjdf_expurgada", "TJ/DF (expurgada)", RATE_DECIMAL, "TJ/DF (expurgada)", "tj_df_expurgada"),
    _spec("tjmg_expurgada", "TJ/MG (expurgada)", VALUE_INDEX, "TJ/MG (expurgada)", "tj_mg_expurgada"),
    _spec("tjro_sem_expurgos", "TJ/RO (s.expurgos)", VALUE_INDEX, "TJRO (s.expurgos)", "tj_ro_sem_expurgos"),
    _spec("tjce_condenatorias_tj_ceara", "TJCE (Condenatórias TJ/Ceará)", VALUE_INDEX),
    _spec("tjdf_nao_expurgada", "TJDF (não expurgada)", RATE_DECIMAL),
    _spec("tjes_tabela_tribunal_just_es", "TJES (Tabela Tribunal Just ES)", RATE_DECIMAL),
    _spec("tjmg_nao_expurgada", "TJMG (não expurgada)", VALUE_INDEX),
    _spec("tjpr_ipca_e_precatorios", "TJPR (IPCA-E / Precatórios)", VALUE_INDEX),
    _spec("tjpr_media_igp_inpc", "TJPR (média IGP/INPC)", RATE_DECIMAL),
    _spec("tjrj_tabela_tribunal_just_rj", "TJRJ (Tabela Tribunal Just RJ)", VALUE_INDEX),
    _spec("tjrs_tabela_tribunal_just_rs_igpm", "TJRS (Tabela Tribunal Just RS-IGPM)", RATE_DECIMAL, "TJRS (Tabela Tribunal Just RS)", "tjrs_igpm"),
    _spec("tjsc_tabela_tribunal_just_sc_icgj", "TJSC - Tabela Tribunal Just SC ICGJ", RATE_DECIMAL),
    _spec("tjsp_inpc_ipca15_lei_14905", "TJSP (INPC/IPCA-15 - Lei 14905)", VALUE_INDEX, "tjsp_inpc_ipca_15_lei_14905"),
    _spec("tjsp_fazenda_publica_precatorios_ate_25_3_15_cnj_303_selic", "TJSP-Fazenda Pública e Precatórios até 25/3/15-CNJ.303 SELIC", VALUE_INDEX),
    _spec("tjsp_precatorios_apos_25_3_15_cnj_303_com_selic", "TJSP-Precatórios após 25/3/15 (CNJ.303 com SELIC)", VALUE_INDEX),
    _spec("tst_debitos_trabalhistas_ipca_e", "TST - Débitos trabalhistas (IPCA-E)", RATE_DECIMAL),
    _spec("tst_debitos_trabalhistas_tr", "TST - Débitos trabalhistas (TR)", RATE_DECIMAL),
)

# Índices conhecidos em documentações de cálculo, mas ausentes do arquivo mensal anexado.
# O projeto não inventa valores: se uma dessas chaves for usada, lança erro claro.
MISSING_INDEX_SPECS: tuple[MissingLocalIndexSpec, ...] = (
    MissingLocalIndexSpec("ipc_ibge_extinto", "IPC-IBGE (extinto)", "mar/1986 a fev/1991"),
    MissingLocalIndexSpec("ipc_r_ibge_extinto", "IPC-R IBGE (extinto)", "jul/1994 a jul/1996"),
    MissingLocalIndexSpec("isn_ibge_extinto", "ISN-IBGE (extinto)", "mar/1991 a abr/1997"),
)

INDEX_SPECS: dict[str, LocalIndexSpec] = {s.key: s for s in _INDEX_SPECS}
ALIAS_TO_KEY: dict[str, str] = {}
for s in _INDEX_SPECS:
    ALIAS_TO_KEY[s.key] = s.key
    ALIAS_TO_KEY[normalize_key(s.key)] = s.key
    ALIAS_TO_KEY[normalize_key(s.column)] = s.key
    for alias in s.aliases:
        ALIAS_TO_KEY[normalize_key(alias)] = s.key
for missing in MISSING_INDEX_SPECS:
    ALIAS_TO_KEY[missing.key] = missing.key
    ALIAS_TO_KEY[normalize_key(missing.label)] = missing.key


def local_index_specs() -> tuple[LocalIndexSpec, ...]:
    """Lista os índices efetivamente disponíveis na planilha mensal.

    Entrada:
        Nenhuma.

    Saída:
        ``tuple[LocalIndexSpec, ...]`` com metadados de cada coluna calculável.
    """
    return _INDEX_SPECS


def local_missing_index_specs() -> tuple[MissingLocalIndexSpec, ...]:
    """Lista índices conhecidos, mas sem dados na planilha mensal.

    Entrada:
        Nenhuma.

    Saída:
        ``tuple[MissingLocalIndexSpec, ...]`` para documentação e erros.
    """
    return MISSING_INDEX_SPECS


def resolve_index_key(key_or_label: str) -> str:
    """Resolve aliases para a chave técnica do índice.

    Entrada:
        ``key_or_label``: chave ou rótulo, como ``"IPCA (IBGE)"``.

    Saída:
        ``str`` com a chave técnica cadastrada, como ``"ipca_ibge"``.
    """
    normalized = normalize_key(key_or_label)
    return ALIAS_TO_KEY.get(normalized, key_or_label)


def get_index_spec(key_or_label: str) -> LocalIndexSpec:
    """Obtém metadados de um índice local.

    Entrada:
        ``key_or_label``: chave ou nome de coluna.

    Saída:
        ``LocalIndexSpec``. Lança ``ValueError`` se o índice não existir ou
        estiver cadastrado sem coluna correspondente.
    """
    key = resolve_index_key(key_or_label)
    if key in INDEX_SPECS:
        return INDEX_SPECS[key]
    missing = {m.key: m for m in MISSING_INDEX_SPECS}
    if key in missing:
        m = missing[key]
        raise ValueError(
            f"O índice '{m.label}' ({m.available_range}) está cadastrado como conhecido, "
            "mas não há coluna correspondente em taxas_mensais.xlsx. "
            "Inclua a coluna na planilha para calcular esse índice sem estimativas."
        )
    disponiveis = ", ".join(sorted(INDEX_SPECS))
    raise ValueError(f"Índice local não encontrado: {key_or_label!r}. Disponíveis: {disponiveis}")


def _decimal_or_none(value: Any) -> Decimal | None:
    """Converte célula de planilha para ``Decimal`` ou ``None``.

    Entrada:
        ``value``: número, texto ou vazio vindo do Excel.

    Saída:
        ``Decimal`` quando há valor; ``None`` para célula vazia.
    """
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    return D(value)


def _month_from_cell(value: Any) -> str | None:
    """Converte célula de competência mensal para ``AAAA-MM``.

    Entrada:
        ``value``: ``datetime``, ``date`` ou texto como ``"08/2024"``.

    Saída:
        ``str`` ``AAAA-MM`` ou ``None`` para célula vazia.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return f"{value.year:04d}-{value.month:02d}"
    if isinstance(value, date):
        return f"{value.year:04d}-{value.month:02d}"
    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d{4}-\d{2}", text):
        return text
    if re.fullmatch(r"\d{2}/\d{4}", text):
        mes, ano = text.split("/")
        return f"{int(ano):04d}-{int(mes):02d}"
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}.*", text):
        return text[:7]
    raise ValueError(f"Competência mensal inválida na planilha: {value!r}")


def _date_from_cell(value: Any) -> date:
    """Converte célula diária da planilha para ``datetime.date``.

    Entrada:
        ``value``: data do Excel, número serial ou texto ``YYYY-MM-DD``/``DD/MM/YYYY``.

    Saída:
        ``datetime.date``.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return from_excel(value).date()
    text = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return datetime.strptime(text, "%Y-%m-%d").date()
    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", text):
        return datetime.strptime(text, "%d/%m/%Y").date()
    raise ValueError(f"Data diária inválida na planilha: {value!r}")


@lru_cache(maxsize=4)
def load_monthly_indices(path: str | None = None) -> pd.DataFrame:
    """Carrega a tabela mensal local em formato largo.

    A primeira coluna vira ``mes``. Os cabeçalhos são stripados para evitar o
    espaço final existente em uma coluna da planilha original.
    """
    file_path = Path(path) if path else _resource_path(MENSAL_XLSX)
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb["indices"] if "indices" in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    raw_headers = next(rows)
    headers = [str(h).strip() if h is not None else "" for h in raw_headers]
    if not headers or normalize_key(headers[0]) not in {"data", "mes"}:
        raise ValueError("A planilha mensal precisa ter a primeira coluna como data/mes.")

    records: list[dict[str, Any]] = []
    for raw in rows:
        mes = _month_from_cell(raw[0] if raw else None)
        if mes is None:
            continue
        rec: dict[str, Any] = {"mes": mes}
        for header, value in zip(headers[1:], raw[1:]):
            if header:
                rec[header] = _decimal_or_none(value)
        records.append(rec)
    if not records:
        raise ValueError("Nenhum dado mensal encontrado em taxas_mensais.xlsx.")
    return pd.DataFrame(records).sort_values("mes").reset_index(drop=True)


@lru_cache(maxsize=128)
def load_index_series(key_or_label: str, path: str | None = None) -> pd.DataFrame:
    """Retorna a série mensal de uma chave local.

    Entrada:
        ``key_or_label``: chave ou rótulo do índice;
        ``path``: caminho opcional para uma planilha mensal alternativa.

    Saída:
        ``DataFrame`` com colunas ``mes`` (``AAAA-MM``) e ``indice`` (``Decimal``).

    Exemplo:
        ``load_index_series("ipca_ibge").head()`` retorna a série mensal do IPCA.
    """
    spec = get_index_spec(key_or_label)
    df = load_monthly_indices(path)
    if spec.column not in df.columns:
        raise ValueError(f"A coluna '{spec.column}' não existe em {MENSAL_XLSX}.")
    out = df[["mes", spec.column]].rename(columns={spec.column: "indice"}).copy()
    out["indice"] = out["indice"].apply(_decimal_or_none)
    out = out.dropna(subset=["indice"]).reset_index(drop=True)
    return out


@lru_cache(maxsize=128)
def local_index_coverage(key_or_label: str, path: str | None = None) -> LocalIndexCoverage:
    """Retorna o intervalo real de uma série, sem datas escritas manualmente.

    A função considera somente células não vazias da coluna selecionada. Assim,
    o catálogo e a validação usam exatamente os dados disponíveis no arquivo de
    cálculo instalado, inclusive após uma atualização das planilhas.
    """
    spec = get_index_spec(key_or_label)
    series = load_index_series(spec.key, path)
    if series.empty:
        raise ValueError(f"A série '{spec.column}' não possui competências preenchidas em {MENSAL_XLSX}.")
    first = str(series["mes"].min())[:7]
    last = str(series["mes"].max())[:7]
    maximum_update = somar_meses(last, 1) if spec.mode == RATE_DECIMAL else last
    return LocalIndexCoverage(
        key=spec.key,
        label=spec.column,
        mode=spec.mode,
        first_competence=first,
        last_competence=last,
        maximum_update_competence=maximum_update,
    )


def load_taxa_legal_mensal_percentual(path: str | None = None) -> pd.DataFrame:
    """Carrega a coluna mensal da Taxa Legal em percentual ao mês.

    A planilha traz a taxa como decimal (0.00605306 = 0,605306%). As funções
    históricas do projeto esperam percentual, então multiplicamos por 100 aqui.
    """
    df = load_monthly_indices(path)
    col = "TAXA LEGAL - art. 406 CC"
    if col not in df.columns:
        raise ValueError(f"A coluna '{col}' não existe em {MENSAL_XLSX}.")
    out = df[["mes", col]].rename(columns={col: "taxa_legal_percentual"}).dropna().copy()
    out["taxa_legal_percentual"] = out["taxa_legal_percentual"].apply(lambda v: D(v) * Decimal("100"))
    return out.reset_index(drop=True)


@lru_cache(maxsize=8)
def load_daily_rate_table(kind: str, path: str | None = None) -> pd.DataFrame:
    """Carrega uma tabela diária local de taxas em decimal ao dia."""
    if path:
        file_path = Path(path)
    elif kind == "selic_ipcae":
        file_path = _resource_path(DIARIA_SELIC_IPCAE_XLSX)
    elif kind == "12_6":
        file_path = _resource_path(DIARIA_12_6_XLSX)
    else:
        raise ValueError("kind deve ser 'selic_ipcae' ou '12_6'.")

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    if normalize_key(headers[0]) != "data":
        raise ValueError("A tabela diária precisa ter a primeira coluna 'data'.")
    value_idx = 1
    records: list[dict[str, Any]] = []
    for raw in rows:
        if not raw or raw[0] is None:
            continue
        value = _decimal_or_none(raw[value_idx] if len(raw) > value_idx else None)
        if value is None:
            continue
        records.append({"data": _date_from_cell(raw[0]), "valor_indice": value})
    if not records:
        raise ValueError(f"Nenhum dado diário encontrado para {kind}.")
    return pd.DataFrame(records).drop_duplicates("data", keep="last").sort_values("data").reset_index(drop=True)


def available_indices() -> pd.DataFrame:
    """Retorna uma tabela de índices locais disponíveis.

    Entrada:
        Nenhuma.

    Saída:
        ``DataFrame`` com ``key``, ``coluna``, ``modo`` e ``aliases``.
    """
    return pd.DataFrame([
        {"key": s.key, "coluna": s.column, "modo": s.mode, "aliases": ", ".join(s.aliases)}
        for s in _INDEX_SPECS
    ])


def missing_indices() -> pd.DataFrame:
    """Retorna índices conhecidos que não possuem coluna na planilha.

    Entrada:
        Nenhuma.

    Saída:
        ``DataFrame`` com ``key``, ``label`` e ``available_range``.
    """
    return pd.DataFrame([m.__dict__ for m in MISSING_INDEX_SPECS])
